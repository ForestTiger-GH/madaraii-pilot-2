from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import coordinate_to_tuple

from .normalization import (
    EN_MONTHS,
    RU_MONTHS,
    infer_unit,
    normalize_text,
    parse_decimal_text,
    parse_period,
    sheet_dimensions,
    stable_dimensions_json,
    territory_type,
    transliterate_ru,
)
from .raw import raw_index
from .registry import SourceSpec


class SourceVariantError(RuntimeError):
    pass


@dataclass(frozen=True)
class PeriodBinding:
    header_row: int
    column: int
    period: str
    frequency: str
    representation: str
    coordinate: str


def _stable_id(prefix: str, *parts: object) -> str:
    payload = "\x1f".join(str(p) for p in parts).encode("utf-8")
    return prefix + hashlib.sha256(payload).hexdigest()[:24]


def _cell_text(cell: Cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _is_period_candidate_value(value: Any) -> bool:
    # Generic discovery intentionally avoids naked numeric years; those occur in data too.
    return isinstance(value, (str, datetime, date))


def _period_header_is_statistical(ws, row_no: int, bindings: list[PeriodBinding]) -> bool:
    """Reject obvious methodology/change-log rows that merely contain several dates."""
    if not bindings:
        return False
    first_period_col = min(b.column for b in bindings)
    context: list[str] = []
    for c in range(1, first_period_col):
        value = ws.cell(row_no, c).value
        if isinstance(value, str) and value.strip():
            context.append(normalize_text(value))
    text = " ".join(context)
    if not text:
        return True
    non_statistical = (
        "методолог" in text
        or "дата изменен" in text
        or "история изменен" in text
        or "revision date" in text
        or "methodolog" in text
    )
    return not non_statistical


def discover_period_bindings(ws, source_id: str) -> list[PeriodBinding]:
    by_row: dict[int, list[PeriodBinding]] = defaultdict(list)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or not _is_period_candidate_value(cell.value):
                continue
            parsed = parse_period(cell.value)
            if parsed:
                period, freq, rep = parsed
                by_row[cell.row].append(
                    PeriodBinding(cell.row, cell.column, period, freq, rep, cell.coordinate)
                )

    accepted: list[PeriodBinding] = []
    for row_no, bindings in by_row.items():
        if len(bindings) >= 3 and _period_header_is_statistical(ws, row_no, bindings):
            accepted.extend(bindings)

    if source_id == "exchange_rate":
        accepted.extend(_discover_exchange_periods(ws))

    uniq: dict[tuple[int, int], PeriodBinding] = {}
    for binding in accepted:
        uniq[(binding.header_row, binding.column)] = binding
    return sorted(uniq.values(), key=lambda b: (b.header_row, b.column))


def _discover_exchange_periods(ws) -> list[PeriodBinding]:
    max_header = min(ws.max_row, 12)
    max_col = ws.max_column
    years_by_row: dict[int, dict[int, int]] = {}

    for r in range(1, max_header + 1):
        years: dict[int, int] = {}
        current: int | None = None
        for c in range(1, max_col + 1):
            value = ws.cell(r, c).value
            year: int | None = None
            if isinstance(value, int) and 1900 <= value <= 2200:
                year = value
            elif isinstance(value, float) and value.is_integer() and 1900 <= value <= 2200:
                year = int(value)
            elif isinstance(value, str) and re.fullmatch(r"\s*(19|20|21)\d{2}\s*", value):
                year = int(value.strip())
            if year:
                current = year
            if current:
                years[c] = current
        if years:
            years_by_row[r] = years

    out: list[PeriodBinding] = []
    for r in range(1, max_header + 1):
        month_or_quarter: list[tuple[int, str, int | None, str]] = []
        for c in range(1, max_col + 1):
            value = ws.cell(r, c).value
            if not isinstance(value, str):
                continue
            token = normalize_text(value).replace(".", "").strip()
            month = RU_MONTHS.get(token) or EN_MONTHS.get(token)
            if month:
                month_or_quarter.append((c, "monthly", month, "exchange_composite"))
                continue
            match = re.fullmatch(r"(?:([iv]+)|([1-4]))\s*(?:кв(?:арт(?:ал)?)?|q)", token)
            if match:
                roman, digit = match.groups()
                quarter = int(digit) if digit else {"i": 1, "ii": 2, "iii": 3, "iv": 4}.get(roman)
                if quarter:
                    month_or_quarter.append((c, "quarterly", quarter, "exchange_composite"))
        if len(month_or_quarter) < 3:
            continue
        prior_year_rows = [years for rr, years in years_by_row.items() if rr < r]
        if not prior_year_rows:
            continue
        years = prior_year_rows[-1]
        for c, freq, token_num, rep in month_or_quarter:
            year = years.get(c)
            if not year:
                continue
            if freq == "monthly":
                period = date(year, int(token_num), 1).isoformat()
            else:
                period = date(year, 1 + 3 * (int(token_num) - 1), 1).isoformat()
            out.append(PeriodBinding(r, c, period, freq, rep, ws.cell(r, c).coordinate))
    return out


def _raw_decimal(raw: dict[str, object], openpyxl_value: Any) -> Decimal | None:
    raw_type = str(raw.get("ooxml_type", ""))
    lexical = str(raw.get("value_lexical", ""))
    resolved = str(raw.get("text_resolved", ""))
    if raw_type in {"n", "str"} and lexical:
        parsed = parse_decimal_text(lexical)
        if parsed is not None:
            return parsed
    if raw_type in {"s", "inlineStr"} and resolved:
        parsed = parse_decimal_text(resolved)
        if parsed is not None:
            return parsed
    return parse_decimal_text(openpyxl_value)


def _raw_value_exact(raw: dict[str, object], openpyxl_value: Any) -> str:
    lexical = str(raw.get("value_lexical", ""))
    if lexical != "":
        return lexical
    resolved = str(raw.get("text_resolved", ""))
    if resolved != "":
        return resolved
    return "" if openpyxl_value is None else str(openpyxl_value)


def _first_title(ws) -> str:
    for r in range(1, min(ws.max_row, 8) + 1):
        for c in range(1, min(ws.max_column, 8) + 1):
            value = ws.cell(r, c).value
            if isinstance(value, str) and value.strip():
                text = " ".join(value.split())
                if len(text) >= 12:
                    return text
    return ws.title


def _row_side_values(ws, row: int, first_period_col: int) -> list[str]:
    values: list[str] = []
    for c in range(1, first_period_col):
        value = ws.cell(row, c).value
        if value is None or isinstance(value, (datetime, date)):
            continue
        text = " ".join(str(value).split())
        if text:
            values.append(text)
    return values


def _row_label(ws, row: int, first_period_col: int) -> str:
    values = _row_side_values(ws, row, first_period_col)
    if not values:
        return f"Строка {row}"
    return " | ".join(values)


def _metadata_row_label(label: str) -> bool:
    text = normalize_text(label)
    return bool(
        re.match(r"^(?:источник|примечани|сноска|техническ(?:ое|ая)? примечани)", text)
        or "источник:" in text
    )


def _readable_ru(label: str) -> str:
    value = " ".join(label.split())
    value = re.sub(r",?\s*(?:млн|млрд|тыс)\.?\s*руб\.?\s*$", "", value, flags=re.I)
    value = re.sub(r",?\s*%\s*$", "", value)
    return value.strip(" ,;:") or label


_COMMON_EN = [
    (r"^итого$", "Total"),
    (r"^всего:?$", "Total"),
    (r"просроченн(?:ая|ой) задолженность", "Overdue debt"),
    (r"задолженность", "Debt"),
    (r"объем кредитов", "Loan volume"),
    (r"количество кредитов", "Number of loans"),
    (r"количество заемщиков", "Number of borrowers"),
    (r"финансовые активы", "Financial assets"),
    (r"обязательства", "Liabilities"),
    (r"денежн(?:ая|ые) масса|денежные агрегаты", "Monetary aggregates"),
    (r"депозиты", "Deposits"),
    (r"кредиты и займы", "Loans"),
    (r"кредиты", "Loans"),
    (r"наличная валюта", "Currency"),
    (r"иностранн(?:ая|ой) валюта", "Foreign currency"),
    (r"национальн(?:ая|ой) валюта", "National currency"),
    (r"долговые ценные бумаги", "Debt securities"),
    (r"краткосрочн", "Short-term"),
    (r"долгосрочн", "Long-term"),
]


def _project_english(label: str, fallback_context: str) -> tuple[str, str]:
    normalized = normalize_text(label)
    for pattern, english in _COMMON_EN:
        if re.search(pattern, normalized, re.I):
            translit = transliterate_ru(_readable_ru(label))
            if len(translit) > len(english) + 12:
                return f"{english} — {translit}", "project_translation"
            return english, "project_translation"
    return f"{fallback_context}: {transliterate_ru(_readable_ru(label))}", "project_translation"


def _embedded_dimensions(label: str, source_id: str) -> dict[str, str]:
    text = normalize_text(label)
    dims: dict[str, str] = {}
    if "иностранн" in text and "валют" in text:
        dims["currency_category"] = "foreign_currency"
    elif "национальн" in text and "валют" in text:
        dims["currency_category"] = "national_currency"
    elif "в рублях" in text or "рублев" in text:
        dims["currency_category"] = "rubles"

    if re.search(r"\bнепросроч", text) or re.search(r"\bбез\b[^\n]{0,40}\bпросроч", text):
        dims["overdue"] = "false"
    elif "просроч" in text:
        dims["overdue"] = "true"

    up_to = "до 1 года" in text or "до одного года" in text
    over = "свыше 1 года" in text or "более 1 года" in text
    if up_to and not over:
        dims["maturity_bucket"] = "up_to_1_year"
    elif over and not up_to:
        dims["maturity_bucket"] = "over_1_year"

    if "остаточн" in text and "срок" in text:
        dims["maturity_basis"] = "remaining"
    elif "первоначальн" in text and "срок" in text:
        dims["maturity_basis"] = "original"
    if source_id.startswith("sme_"):
        dims["sme_scope"] = "sme"
    if source_id == "sme_debt_subj_i":
        dims["entrepreneur_scope"] = "individual_entrepreneurs"
    if source_id == "dep_ind_no_escrow":
        dims["escrow_coverage"] = "excluded"
    if source_id == "escrow_accounts":
        dims["escrow_coverage"] = "escrow_only"
    return dims


def _period_role(spec: SourceSpec, label: str, sheet: str) -> str:
    text = normalize_text(label + " " + sheet)
    if "методолог" in text or "изменение отсутств" in text or "без изменен" in text:
        return spec.period_role
    if "темп прироста" in text or "прирост" in text or "изменение" in text:
        return "published_change"
    if "за месяц" in text or "за отчетный период" in text or "операци" in text or "объем кредитов, предоставленных" in text:
        return "flow"
    if "задолж" in text or "остат" in text or "баланс" in text or "по состоянию" in text:
        return "stock"
    return spec.period_role


def _activity_code(label: str) -> str:
    match = re.match(r"\s*([A-ZА-Я]?\d{1,2}(?:\.\d+)*)\b", label)
    return match.group(1) if match else ""


def _concept(
    spec: SourceSpec,
    *,
    sheet: str,
    row: int,
    label: str,
    title: str,
    row_axis: str,
    unit: str | None = None,
    scale: int | None = None,
) -> dict[str, str]:
    if row_axis in {"region", "activity"}:
        measure_label = _readable_ru(title)
        local_key = "|".join(
            (
                "dataset_measure",
                normalize_text(measure_label),
                unit or "source_defined",
                "" if scale is None else str(scale),
            )
        )
        label_source = title
        name_ru = measure_label if normalize_text(measure_label) != normalize_text(sheet) else spec.name_ru
        if normalize_text(measure_label) == normalize_text(spec.name_ru):
            name_en = spec.name_en
        else:
            name_en, _ = _project_english(measure_label, spec.name_en)
        translation_status = "project_translation"
    else:
        # Row position is presentation geometry, not durable source-local identity.
        local_key = "|".join((normalize_text(sheet), normalize_text(title), normalize_text(label)))
        label_source = label
        name_ru = _readable_ru(label)
        name_en, translation_status = _project_english(label, spec.name_en)
    concept_id = _stable_id("sc_", spec.source_id, local_key)
    return {
        "source_concept_id": concept_id,
        "source_id": spec.source_id,
        "source_local_key": local_key,
        "label_ru_source": label_source,
        "label_ru_normalized": normalize_text(label_source),
        "name_ru": name_ru,
        "name_en": name_en,
        "translation_status": translation_status,
        "row_axis": row_axis,
        "source_context": title,
    }


def _dimension_member(dimension: str, source_value: str, *, classification: str | None = None) -> dict[str, str]:
    norm = normalize_text(source_value)
    member_id = _stable_id("dm_", dimension, classification or "", norm)
    if dimension == "region":
        english = transliterate_ru(source_value)
        status = "transliteration_only"
    elif dimension == "activity":
        code = _activity_code(source_value)
        prefix = classification or "activity"
        english = f"{prefix} {code}: {transliterate_ru(source_value)}" if code else f"{prefix}: {transliterate_ru(source_value)}"
        status = "project_translation"
    else:
        english = source_value
        status = "project_translation"
    return {
        "dimension_member_id": member_id,
        "dimension": dimension,
        "classification": classification or "",
        "source_value_ru": source_value,
        "name_ru": source_value,
        "name_en": english,
        "translation_status": status,
    }


def parse_source(
    xlsx_path: str | Path,
    *,
    spec: SourceSpec,
    source_revision_id: str,
    file_sha256: str,
    raw_cells: list[dict[str, object]],
) -> tuple[list[dict[str, object]], list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], dict[str, object]]:
    """Return observations, concepts, dimension members, dispositions, source diagnostics."""
    wb = load_workbook(xlsx_path, read_only=False, data_only=False, keep_links=True)
    rindex = raw_index(raw_cells)
    dispositions: dict[str, dict[str, str]] = {
        str(row["raw_cell_id"]): {
            "raw_cell_id": str(row["raw_cell_id"]),
            "source_id": spec.source_id,
            "source_revision_id": source_revision_id,
            "role": "source_cell",
            "reason": "retained_nonempty_source_cell",
        }
        for row in raw_cells
    }
    observations: list[dict[str, object]] = []
    concepts: dict[str, dict[str, str]] = {}
    members: dict[str, dict[str, str]] = {}
    sheet_diags: list[dict[str, object]] = []

    for ws in wb.worksheets:
        title = _first_title(ws)
        bindings = discover_period_bindings(ws, spec.source_id)
        by_header: dict[int, dict[int, PeriodBinding]] = defaultdict(dict)
        for binding in bindings:
            by_header[binding.header_row][binding.column] = binding
            raw = rindex.get((ws.title, binding.coordinate))
            if raw:
                dispositions[str(raw["raw_cell_id"])].update(
                    {"role": "period_key", "reason": binding.representation}
                )

        header_rows = sorted(by_header)
        sheet_dims = sheet_dimensions(ws.title, spec.source_id)
        metadata_sheet = any(
            token in normalize_text(ws.title)
            for token in ("методолог", "metadata", "метадан")
        )
        obs_before = len(observations)
        numeric_candidates = 0

        raw_by_coord = {
            coord: row for (sheet, coord), row in rindex.items() if sheet == ws.title
        }

        for coord, raw in raw_by_coord.items():
            row_no, col_no = coordinate_to_tuple(coord)
            open_value = ws[coord].value
            dec = _raw_decimal(raw, open_value)
            if dec is None:
                continue
            numeric_candidates += 1

            candidate_headers = [
                hr for hr in header_rows if hr < row_no and col_no in by_header[hr]
            ]
            if not candidate_headers:
                role = "non_observation_numeric"
                reason = "no_period_binding"
                if col_no <= 5:
                    role, reason = "hierarchy_or_header_code", "numeric_before_period_axis"
                dispositions[str(raw["raw_cell_id"])].update({"role": role, "reason": reason})
                continue

            header_row = candidate_headers[-1]
            binding = by_header[header_row][col_no]
            first_period_col = min(by_header[header_row])
            if row_no <= header_row:
                continue

            row_label = _row_label(ws, row_no, first_period_col)
            if metadata_sheet or _metadata_row_label(row_label):
                dispositions[str(raw["raw_cell_id"])].update(
                    {"role": "source_metadata_numeric", "reason": "metadata_or_note_row"}
                )
                continue

            dims: dict[str, str] = dict(sheet_dims)
            dims.update(_embedded_dimensions(row_label, spec.source_id))
            if spec.row_axis == "region":
                dims["region"] = row_label
                dims["region_type"] = territory_type(row_label)
                member = _dimension_member("region", row_label)
                members[member["dimension_member_id"]] = member
            elif spec.row_axis == "activity":
                dims["activity"] = row_label
                dims["classification"] = spec.classification or "source_activity"
                code = _activity_code(row_label)
                if code:
                    dims["activity_code"] = code
                member = _dimension_member("activity", row_label, classification=spec.classification)
                members[member["dimension_member_id"]] = member

            unit, scale = infer_unit(row_label, title)
            concept = _concept(
                spec,
                sheet=ws.title,
                row=row_no,
                label=row_label,
                title=title,
                row_axis=spec.row_axis,
                unit=unit,
                scale=scale,
            )
            concepts[concept["source_concept_id"]] = concept
            obs_id = _stable_id("ob_", source_revision_id, ws.title, coord)
            observations.append(
                {
                    "observation_id": obs_id,
                    "source_id": spec.source_id,
                    "source_revision_id": source_revision_id,
                    "file_sha256": file_sha256,
                    "source_concept_id": concept["source_concept_id"],
                    "period": binding.period,
                    "frequency": "monthly" if binding.frequency == "monthly_or_point" else binding.frequency,
                    "period_representation": binding.representation,
                    "period_role": _period_role(spec, row_label, ws.title),
                    "value_exact": _raw_value_exact(raw, open_value),
                    "value_kind": "formula_cached" if str(raw.get("formula", "")) else (
                        "numeric_text" if str(raw.get("ooxml_type")) in {"s", "inlineStr", "str"} else "numeric"
                    ),
                    "unit": unit or "source_defined",
                    "scale": scale if scale is not None else "",
                    "dimensions_json": stable_dimensions_json(dims),
                    "sheet_exact": ws.title,
                    "cell_coordinate": coord,
                    "raw_cell_id": str(raw["raw_cell_id"]),
                    "source_row": row_no,
                    "source_row_label": row_label,
                }
            )
            dispositions[str(raw["raw_cell_id"])].update(
                {"role": "observation_value", "reason": "period_bound_numeric_value"}
            )

        obs_count = len(observations) - obs_before
        sheet_diags.append(
            {
                "sheet": ws.title,
                "title": title,
                "metadata_sheet": metadata_sheet,
                "period_bindings": len(bindings),
                "period_header_rows": header_rows,
                "numeric_candidates": numeric_candidates,
                "observations": obs_count,
            }
        )
        if numeric_candidates >= 10 and not metadata_sheet and obs_count == 0:
            wb.close()
            raise SourceVariantError(
                f"{spec.source_id}/{ws.title}: {numeric_candidates} numeric candidates but no semantic observations; period/header contract unrecognized"
            )

    wb.close()
    return observations, list(concepts.values()), list(members.values()), list(dispositions.values()), {
        "source_id": spec.source_id,
        "source_revision_id": source_revision_id,
        "sheets": sheet_diags,
        "observation_count": len(observations),
        "concept_count": len(concepts),
        "dimension_member_count": len(members),
    }
