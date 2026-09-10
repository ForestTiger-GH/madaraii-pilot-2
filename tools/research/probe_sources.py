from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import tempfile
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


@dataclass(frozen=True)
class Source:
    source_id: str
    url: str
    registry_label: str
    family_hint: str


SOURCES: tuple[Source, ...] = (
    Source("mortgage_debt_ind", "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_05_Debt_ind.xlsx", "Кредиты физлиц: задолженность", "mortgage_and_household_loans"),
    Source("mortgage_debt", "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_14_Debt_mortgage.xlsx", "Ипотека: задолженность", "mortgage_and_household_loans"),
    Source("mortgage_scpa_debt", "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_18_Debt_scpa_mortgage.xlsx", "Ипотека: задолженность по ДДУ", "mortgage_and_household_loans"),
    Source("mortgage_ihc", "https://www.cbr.ru/vfs/statistics/banksector/mortgage/02_41_Mortgage_ihc.xlsx", "Ипотека: жильё в строящихся домах", "mortgage_and_household_loans"),
    Source("mortgage_full", "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_02_Mortgage.xlsx", "Ипотека: полный ряд", "mortgage_and_household_loans"),
    Source("mortgage_scpa_full", "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_03_Scpa_mortgage.xlsx", "Ипотека: ДДУ полный ряд", "mortgage_and_household_loans"),
    Source("corp_new_loans_a", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_01_A_New_loans_corp_by_activity.xlsx", "Корпорации: новые кредиты A", "corporate_and_sme_loans"),
    Source("corp_new_loans_c", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_01_C_New_loans_corp_by_activity.xlsx", "Корпорации: новые кредиты C", "corporate_and_sme_loans"),
    Source("corp_debt_a", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_02_A_Debt_corp_by_activity.xlsx", "Корпорации: долг A", "corporate_and_sme_loans"),
    Source("corp_debt_c", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_02_C_Debt_corp_by_activity.xlsx", "Корпорации: долг C", "corporate_and_sme_loans"),
    Source("sme_debt", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_11_Debt_sme.xlsx", "МСП: долг", "corporate_and_sme_loans"),
    Source("sme_debt_activity", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_11_F_Debt_sme_by_activity.xlsx", "МСП: долг по видам деятельности", "corporate_and_sme_loans"),
    Source("sme_debt_subj_f", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_13_F_Debt_sme_subj.xlsx", "МСП: долг по субъектам F", "corporate_and_sme_loans"),
    Source("sme_debt_subj_i", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_13_I_Debt_sme_subj.xlsx", "МСП: долг по субъектам I", "corporate_and_sme_loans"),
    Source("corp_debt_subj", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_05_D_Debt_subj.xlsx", "Корпорации: долг по субъектам", "corporate_and_sme_loans"),
    Source("debt_structure_benchmark_rate", "https://cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/02_03_Debt_structure_by_benchmark_interest_rate_type.xlsx", "Структура долга по типу бенчмарк-ставки", "corporate_and_sme_loans"),
    Source("sme_borrowers_info", "https://www.cbr.ru/vfs/statistics/banksector/loans_to_corporations/02_02_SME_Borrowers_info.xlsx", "МСП: сведения о заёмщиках", "corporate_and_sme_loans"),
    Source("sme_by_activity", "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/01_10_F_New_loans_sme_by_activity.xlsx", "МСП: выдачи по видам деятельности", "corporate_and_sme_loans"),
    Source("debt_securities", "https://www.cbr.ru/vfs/statistics/debt_securities/66-debt_securities.xlsx", "Долговые ценные бумаги", "debt_securities"),
    Source("funds_all", "https://www.cbr.ru/vfs/statistics/BankSector/Borrowings/02_01_Funds_all.xlsx", "Средства совокупные", "borrowings"),
    Source("funds_clients", "https://www.cbr.ru/vfs/statistics/BankSector/Borrowings/02_02_Funds_clients.xlsx", "Средства клиентов", "borrowings"),
    Source("funds_org", "https://www.cbr.ru/vfs/statistics/BankSector/Borrowings/02_04_Funds_org.xlsx", "Средства организаций", "borrowings"),
    Source("dep_corp", "https://www.cbr.ru/vfs/statistics/BankSector/Borrowings/02_05_Dep_corp.xlsx", "Депозиты корпораций", "borrowings"),
    Source("dep_ind", "https://www.cbr.ru/vfs/statistics/BankSector/Borrowings/02_06_Dep_ind.xlsx", "Депозиты физлиц", "borrowings"),
    Source("dep_ind_no_escrow", "https://www.cbr.ru/vfs/statistics/banksector/borrowings/02_27_Dep_ind_excluding_escrow.xlsx", "Депозиты физлиц без escrow", "borrowings"),
    Source("dep_entrepreneur", "https://www.cbr.ru/vfs/statistics/BankSector/Borrowings/02_07_Dep_enterpreneur.xlsx", "Депозиты ИП", "borrowings"),
    Source("escrow_accounts", "https://www.cbr.ru/vfs/statistics/banksector/borrowings/02_28_Escrow_accounts.xlsx", "Счета эскроу", "borrowings"),
    Source("budget_all", "https://www.cbr.ru/vfs/statistics/banksector/borrowings/02_29_Budget_all.xlsx", "Средства бюджетов", "borrowings"),
    Source("households_bm", "https://cbr.ru/vfs/statistics/households/households_bm.xlsx", "Домохозяйства BM", "households"),
    Source("households_om", "https://cbr.ru/vfs/statistics/households/households_om.xlsx", "Домохозяйства OM", "households"),
    Source("monetary_agg", "https://www.cbr.ru/vfs/statistics/credit_statistics/monetary_agg.xlsx", "Денежные агрегаты", "monetary_financial_statistics"),
    Source("survey_cb", "https://www.cbr.ru/vfs/statistics/credit_statistics/survey/survey_cb.xlsx", "Обзор центрального банка", "monetary_financial_statistics"),
    Source("balance_odc", "https://www.cbr.ru/vfs/statistics/credit_statistics/survey/balance_odc.xlsx", "Баланс кредитных организаций", "monetary_financial_statistics"),
    Source("survey_odc", "https://www.cbr.ru/vfs/statistics/credit_statistics/survey/survey_odc.xlsx", "Обзор кредитных организаций", "monetary_financial_statistics"),
    Source("survey_dc_new", "https://www.cbr.ru/vfs/statistics/credit_statistics/survey/survey_dc_new.xlsx", "Обзор банковской системы", "monetary_financial_statistics"),
    Source("annex_survey_dc", "https://www.cbr.ru/vfs/statistics/credit_statistics/survey/annex_survey_dc.xlsx", "Приложение к Обзору банковской системы", "monetary_financial_statistics"),
    Source("debt_new", "https://www.cbr.ru/vfs/statistics/credit_statistics/debt/debt_new.xlsx", "Внешний долг Российской Федерации", "external_debt"),
    Source("debt_maturity", "https://www.cbr.ru/vfs/statistics/credit_statistics/debt/debt_maturity.xlsx", "Внешний долг Российской Федерации по срокам погашения и финансовым инструментам", "external_debt"),
    Source("debt_cur-mat_new", "https://www.cbr.ru/vfs/statistics/credit_statistics/debt/debt_cur-mat_new.xlsx", "Внешний долг Российской Федерации в национальной и иностранной валютах", "external_debt"),
    Source("obs_table_20s", "https://www.cbr.ru/Content/Document/File/115862/obs_tabl20%D1%81.xlsx", "Таблица obs_tabl20с", "other"),
    Source("exchange_rate", "https://cbr.ru/vfs/statistics/credit_statistics/ex_rate_ind/exchange_rate.xlsx", "Курс валют", "exchange_rates"),
)

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.0.0 Safari/537.36"
}


def normalized_label(value: str) -> str:
    s = unicodedata.normalize("NFKC", value).replace("\u00a0", " ").replace("\u200b", "")
    s = re.sub(r"\s+", " ", s).strip().casefold()
    return s


def presentation_signature(value: str) -> str:
    s = normalized_label(value)
    s = re.sub(r"[\s\.,;:!?'\"`’‘“”\-_–—/\\()\[\]{}]+", "", s)
    return s


def json_value(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return {"kind": "date", "value": v.isoformat()}
    if isinstance(v, bytes):
        return {"kind": "bytes", "value": v.hex()}
    if isinstance(v, float):
        return {"kind": "float", "value": repr(v)}
    if isinstance(v, int):
        return {"kind": "int", "value": str(v)}
    if isinstance(v, bool):
        return {"kind": "bool", "value": v}
    if v is None:
        return None
    return str(v)


def make_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=4,
        connect=4,
        read=4,
        status=4,
        backoff_factor=1.0,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET", "HEAD"}),
        respect_retry_after_header=True,
    )
    session.mount("https://", HTTPAdapter(max_retries=retry))
    session.mount("http://", HTTPAdapter(max_retries=retry))
    session.headers.update(HEADERS)
    return session


def inspect_workbook(path: Path, source: Source) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    result: dict[str, Any] = {
        "source": asdict(source),
        "sheet_names": wb.sheetnames,
        "defined_names": [str(x) for x in wb.defined_names.values()],
        "sheets": [],
    }
    sheet_index: list[dict[str, Any]] = []
    label_index: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        cells = [c for c in ws._cells.values() if c.value is not None]
        cells.sort(key=lambda c: (c.row, c.column))
        counts = Counter()
        number_formats = Counter()
        text_cells: list[dict[str, Any]] = []
        formula_cells: list[dict[str, Any]] = []
        numeric_examples: list[dict[str, Any]] = []
        date_examples: list[dict[str, Any]] = []
        rows: dict[int, Counter] = {}
        cols: dict[int, Counter] = {}

        min_row = min((c.row for c in cells), default=None)
        max_row = max((c.row for c in cells), default=None)
        min_col = min((c.column for c in cells), default=None)
        max_col = max((c.column for c in cells), default=None)

        for c in cells:
            v = c.value
            number_formats[c.number_format or ""] += 1
            rowc = rows.setdefault(c.row, Counter())
            colc = cols.setdefault(c.column, Counter())

            if c.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                kind = "formula"
                if len(formula_cells) < 2000:
                    formula_cells.append({"coordinate": c.coordinate, "formula": str(v)})
            elif isinstance(v, bool):
                kind = "bool"
            elif isinstance(v, (datetime, date)) or c.is_date:
                kind = "date"
                if len(date_examples) < 80:
                    date_examples.append({"coordinate": c.coordinate, "value": json_value(v), "number_format": c.number_format})
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                kind = "numeric"
                if len(numeric_examples) < 80:
                    numeric_examples.append({"coordinate": c.coordinate, "value": json_value(v), "number_format": c.number_format})
            elif isinstance(v, str):
                kind = "text"
                text = v.strip("\r\n")
                rec = {
                    "coordinate": c.coordinate,
                    "row": c.row,
                    "column": c.column,
                    "value": text,
                    "normalized": normalized_label(text),
                    "presentation_signature": presentation_signature(text),
                }
                if len(text_cells) < 50000:
                    text_cells.append(rec)
                label_index.append({
                    "source_id": source.source_id,
                    "family_hint": source.family_hint,
                    "sheet": ws.title,
                    "sheet_normalized": normalized_label(ws.title),
                    "sheet_signature": presentation_signature(ws.title),
                    **rec,
                })
            else:
                kind = "other"

            counts[kind] += 1
            rowc[kind] += 1
            colc[kind] += 1

        row_profile = [
            {"row": r, "nonempty": sum(cnt.values()), **dict(cnt)}
            for r, cnt in sorted(rows.items())
        ]
        col_profile = [
            {"column": c, "nonempty": sum(cnt.values()), **dict(cnt)}
            for c, cnt in sorted(cols.items())
        ]

        sheet_rec = {
            "title": ws.title,
            "title_normalized": normalized_label(ws.title),
            "title_presentation_signature": presentation_signature(ws.title),
            "state": ws.sheet_state,
            "declared_max_row": ws.max_row,
            "declared_max_column": ws.max_column,
            "actual_nonempty_bounds": {"min_row": min_row, "max_row": max_row, "min_column": min_col, "max_column": max_col},
            "nonempty_cells": len(cells),
            "cell_type_counts": dict(counts),
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "number_formats_top": number_formats.most_common(30),
            "text_cells": text_cells,
            "text_cells_truncated": counts.get("text", 0) > len(text_cells),
            "formula_cells": formula_cells,
            "formula_cells_truncated": counts.get("formula", 0) > len(formula_cells),
            "numeric_examples": numeric_examples,
            "date_examples": date_examples,
            "row_profile": row_profile,
            "column_profile": col_profile,
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter_ref": ws.auto_filter.ref,
        }
        result["sheets"].append(sheet_rec)

        sheet_index.append({
            "source_id": source.source_id,
            "family_hint": source.family_hint,
            "sheet": ws.title,
            "sheet_normalized": normalized_label(ws.title),
            "sheet_signature": presentation_signature(ws.title),
            "state": ws.sheet_state,
            "nonempty_cells": len(cells),
            "text_cells": counts.get("text", 0),
            "numeric_cells": counts.get("numeric", 0),
            "formula_cells": counts.get("formula", 0),
            "date_cells": counts.get("date", 0),
            "min_row": min_row,
            "max_row": max_row,
            "min_column": min_col,
            "max_column": max_col,
            "merged_ranges": len(ws.merged_cells.ranges),
        })

    wb.close()
    return result, sheet_index, label_index


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                fields.append(key)
                seen.add(key)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "_mw/research/evidence/source-probe")
    out.mkdir(parents=True, exist_ok=True)
    session = make_session()
    manifest: list[dict[str, Any]] = []
    all_sheet_index: list[dict[str, Any]] = []
    all_label_index: list[dict[str, Any]] = []

    with tempfile.TemporaryDirectory(prefix="cbr-probe-") as td:
        tmp = Path(td)
        for i, source in enumerate(SOURCES, 1):
            print(f"[{i:02d}/{len(SOURCES)}] {source.source_id}", flush=True)
            rec: dict[str, Any] = {"source": asdict(source), "status": "pending"}
            try:
                response = session.get(source.url, timeout=(20, 180), allow_redirects=True)
                rec["http_status"] = response.status_code
                rec["final_url"] = response.url
                rec["content_type"] = response.headers.get("Content-Type")
                rec["etag"] = response.headers.get("ETag")
                rec["last_modified"] = response.headers.get("Last-Modified")
                rec["content_length_header"] = response.headers.get("Content-Length")
                response.raise_for_status()
                content = response.content
                rec["bytes"] = len(content)
                rec["sha256"] = hashlib.sha256(content).hexdigest()
                rec["zip_magic"] = content[:2] == b"PK"
                if not rec["zip_magic"]:
                    raise ValueError(f"Downloaded content is not an OOXML ZIP; first bytes={content[:16]!r}")
                local = tmp / f"{source.source_id}.xlsx"
                local.write_bytes(content)
                workbook, sheet_rows, label_rows = inspect_workbook(local, source)
                workbook.update({
                    "acquisition": {
                        "requested_url": source.url,
                        "final_url": response.url,
                        "http_status": response.status_code,
                        "bytes": len(content),
                        "sha256": rec["sha256"],
                        "content_type": rec["content_type"],
                        "etag": rec["etag"],
                        "last_modified": rec["last_modified"],
                    }
                })
                (out / f"{source.source_id}.json").write_text(
                    json.dumps(workbook, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                rec["status"] = "inspected"
                rec["sheets"] = len(workbook["sheets"])
                rec["sheet_names"] = workbook["sheet_names"]
                rec["nonempty_cells"] = sum(s["nonempty_cells"] for s in workbook["sheets"])
                rec["text_cells"] = sum(s["cell_type_counts"].get("text", 0) for s in workbook["sheets"])
                rec["numeric_cells"] = sum(s["cell_type_counts"].get("numeric", 0) for s in workbook["sheets"])
                rec["formula_cells"] = sum(s["cell_type_counts"].get("formula", 0) for s in workbook["sheets"])
                all_sheet_index.extend(sheet_rows)
                all_label_index.extend(label_rows)
            except Exception as exc:  # Research evidence must preserve every failure.
                rec["status"] = "failed"
                rec["error_type"] = type(exc).__name__
                rec["error"] = str(exc)
            manifest.append(rec)

    (out / "MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(out / "SHEET_INDEX.csv", all_sheet_index)
    write_csv(out / "LABEL_INDEX.csv", all_label_index)

    inspected = [r for r in manifest if r["status"] == "inspected"]
    failed = [r for r in manifest if r["status"] != "inspected"]
    source_family_counts = Counter(r["source"]["family_hint"] for r in manifest)
    sheet_signature_counts = Counter(r["sheet_signature"] for r in all_sheet_index)
    repeated_sheet_signatures = [(k, v) for k, v in sheet_signature_counts.most_common() if v > 1]

    lines = [
        "# RT-CBR-001 automated source probe evidence",
        "",
        "This is mechanical research evidence, not a semantic normalization decision.",
        "",
        f"- Registry sources: **{len(manifest)}**",
        f"- Successfully acquired and inspected: **{len(inspected)}**",
        f"- Failed: **{len(failed)}**",
        f"- Material sheets inspected: **{len(all_sheet_index)}**",
        f"- Text cells indexed: **{len(all_label_index)}**",
        "",
        "## Registry family hints",
        "",
    ]
    for family, count in sorted(source_family_counts.items()):
        lines.append(f"- `{family}`: {count} files")
    lines += ["", "## Source dispositions", "", "| Source | Status | Sheets | Bytes | SHA-256 |", "| --- | --- | ---: | ---: | --- |"]
    for r in manifest:
        lines.append(
            f"| `{r['source']['source_id']}` | {r['status']} | {r.get('sheets','')} | {r.get('bytes','')} | `{r.get('sha256','')}` |"
        )
    if failed:
        lines += ["", "## Acquisition/inspection failures", ""]
        for r in failed:
            lines.append(f"- `{r['source']['source_id']}` — {r.get('error_type')}: {r.get('error')}")
    lines += ["", "## Repeated presentation-normalized sheet-name signatures", ""]
    for sig, count in repeated_sheet_signatures[:100]:
        lines.append(f"- `{sig}` — {count} sheets")
    lines += [
        "",
        "## Evidence files",
        "",
        "- `MANIFEST.json` — acquisition/hash/disposition for every registry item.",
        "- `SHEET_INDEX.csv` — one row per workbook sheet with geometry/type counts.",
        "- `LABEL_INDEX.csv` — all observed text labels with exact coordinate plus conservative and punctuation-insensitive presentation signatures.",
        "- `<source_id>.json` — per-workbook sheet details, labels, merged ranges, formulas, number formats, row/column density and representative numeric/date anchors.",
        "",
        "Semantic equivalence must be established by subsequent Research; repeated signatures here are discovery candidates only.",
    ]
    (out / "SUMMARY.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Inspected {len(inspected)}/{len(manifest)} sources; failed={len(failed)}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
