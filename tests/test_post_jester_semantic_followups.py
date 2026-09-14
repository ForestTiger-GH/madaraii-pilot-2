import json

from openpyxl import Workbook

from cbr_unified.processing import parse_source_checked
from cbr_unified.raw import extract_raw_cells, workbook_sha256
from cbr_unified.registry import SourceSpec


def _spec(source_id="synthetic", row_axis="hierarchy", parser="hierarchy"):
    return SourceSpec(
        source_id=source_id,
        url="https://cbr.ru/test.xlsx",
        name_ru="Тест",
        name_en="Test",
        family="test",
        row_axis=row_axis,
        parser=parser,
        period_role="stock",
    )


def _parse(path, spec):
    sha = workbook_sha256(path)
    revision = f"{spec.source_id}@sha256:{sha}"
    raw = extract_raw_cells(
        path,
        source_id=spec.source_id,
        source_revision_id=revision,
        file_sha256=sha,
    )
    return parse_source_checked(
        path,
        spec=spec,
        source_revision_id=revision,
        file_sha256=sha,
        raw_cells=raw,
    )


def test_exchange_axis_measure_headings_split_identical_indicator_labels(tmp_path):
    path = tmp_path / "exchange-like.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Ежемесячные"
    ws["B1"] = "Основные производные показатели динамики обменного курса рубля"
    ws["B2"], ws["C2"], ws["D2"] = 2024, 2024, 2024
    ws["B3"], ws["C3"], ws["D3"] = "Янв", "Фев", "Мар"

    sections = (
        (4, "Индексы обменного курса рубля (в % прироста к декабрю предыдущего года) *", 0.2),
        (6, "Индексы обменного курса рубля (в % прироста к предыдущему периоду) *", -1.2),
        (8, "Индексы обменного курса рубля (в % прироста к соответствующему периоду предыдущего года) *", 4.5),
    )
    for heading_row, heading, base in sections:
        ws.cell(heading_row, 2).value = heading
        data_row = heading_row + 1
        ws.cell(data_row, 1).value = "Индекс номинального эффективного курса рубля к иностранным валютам"
        ws.cell(data_row, 2).value = base
        ws.cell(data_row, 3).value = base + 0.1
        ws.cell(data_row, 4).value = base + 0.2

    wb.save(path)
    wb.close()

    observations, concepts, _, _, diagnostics = _parse(
        path,
        _spec("exchange_rate", "indicator", "exchange"),
    )
    concept_by_id = {row["source_concept_id"]: row for row in concepts}
    observed_contexts = {}
    for source_row in (5, 7, 9):
        concept_id = next(
            row["source_concept_id"]
            for row in observations
            if row["source_row"] == source_row
        )
        observed_contexts[source_row] = concept_by_id[concept_id]["source_context"].lower()

    assert len(set(
        next(row["source_concept_id"] for row in observations if row["source_row"] == source_row)
        for source_row in (5, 7, 9)
    )) == 3
    assert "к декабрю предыдущего года" in observed_contexts[5], observed_contexts
    assert "к предыдущему периоду" in observed_contexts[7], observed_contexts
    assert "к соответствующему периоду предыдущего года" in observed_contexts[9], observed_contexts
    assert diagnostics["exchange_measure_context_rewrites"] == 9


def test_inherited_currency_scope_does_not_leak_into_unique_following_total(tmp_path):
    path = tmp_path / "currency-aggregate-boundary.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "2012-2026"
    ws["A1"] = "Внешний долг Российской Федерации"
    ws["B3"], ws["C3"], ws["D3"] = "01.01.2024", "01.04.2024", "01.07.2024"

    ws["A4"] = "Иностранная валюта"
    ws["A5"] = "Долговые ценные бумаги"
    ws["B5"], ws["C5"], ws["D5"] = 10, 11, 12

    ws["A7"] = "Национальная валюта"
    ws["A8"] = "Долговые ценные бумаги"
    ws["B8"], ws["C8"], ws["D8"] = 20, 21, 22

    ws["A10"] = "Итого по сектору"
    ws["B10"], ws["C10"], ws["D10"] = 30, 32, 34

    wb.save(path)
    wb.close()

    observations, _, _, _, diagnostics = _parse(path, _spec())
    foreign = {
        json.loads(row["dimensions_json"])["currency_category"]
        for row in observations if row["source_row"] == 5
    }
    rubles = {
        json.loads(row["dimensions_json"])["currency_category"]
        for row in observations if row["source_row"] == 8
    }
    aggregate_dims = [
        json.loads(row["dimensions_json"])
        for row in observations if row["source_row"] == 10
    ]
    assert foreign == {"foreign_currency"}
    assert rubles == {"rubles"}
    assert aggregate_dims and all("currency_category" not in dims for dims in aggregate_dims)
    assert diagnostics["structural_scope_dimension_removals"] == 3
