import json

import pytest
from openpyxl import Workbook

from cbr_unified.build import _fingerprints
from cbr_unified.cli import _dims
from cbr_unified.normalization import infer_unit, parse_decimal_text, sheet_dimensions, territory_type
from cbr_unified.raw import extract_raw_cells, workbook_sha256
from cbr_unified.registry import SourceSpec, get_source
from cbr_unified.semantic import (
    _concept,
    _embedded_dimensions,
    _period_role,
    discover_period_bindings,
)
from cbr_unified.processing import parse_source_checked
from cbr_unified.validation import ValidationError, validate_bundle


def _spec(row_axis="indicator"):
    return SourceSpec(
        source_id="synthetic",
        url="https://cbr.ru/test.xlsx",
        name_ru="Тест",
        name_en="Test",
        family="test",
        row_axis=row_axis,
        period_role="stock",
    )


def test_context_inference_respects_negation_specificity_and_ambiguity():
    assert _embedded_dimensions("Непросроченная задолженность", "x")["overdue"] == "false"
    assert sheet_dimensions("Без просроченной задолженности")["overdue"] == "false"
    assert sheet_dimensions("Без права требования")["acquired_claims"] == "excluded"
    assert sheet_dimensions("Несезонно скорректированные данные")["adjustment"] == "original"
    assert infer_unit("Темп прироста, %", "Показатели, млн руб.") == ("percent", 1)
    assert infer_unit("Количество заемщиков, единиц", "Показатели, млн руб.") == ("count", 1)
    assert "maturity_bucket" not in _embedded_dimensions("Итого: до 1 года и свыше 1 года", "x")
    assert territory_type("Российская Федерация без данных по отдельной территории") == "excluding_subregion_aggregate"
    assert _period_role(_spec(), "Задолженность", "Изменение методологии") == "stock"


def test_accounting_parentheses_are_inside_numeric_completeness_boundary():
    assert parse_decimal_text("(10)") == -10
    assert parse_decimal_text("(1 234,50)") == parse_decimal_text("-1234,50")


def test_methodology_date_row_is_not_a_period_axis(tmp_path):
    path = tmp_path / "methodology.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A2"] = "Дата изменения методологии"
    for col, value in zip((2, 3, 4), ("01.01.2024", "01.02.2024", "01.03.2024")):
        ws.cell(2, col).value = value
    ws["A3"] = "Контроль"
    ws["B3"], ws["C3"], ws["D3"] = 1, 2, 3
    wb.save(path)
    wb.close()

    from openpyxl import load_workbook

    loaded = load_workbook(path)
    try:
        assert discover_period_bindings(loaded["Data"], "synthetic") == []
    finally:
        loaded.close()


def test_technical_note_number_is_not_admitted_as_observation(tmp_path):
    path = tmp_path / "notes.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Задолженность, млн руб."
    for col, value in zip((2, 3, 4), ("01.01.2024", "01.02.2024", "01.03.2024")):
        ws.cell(2, col).value = value
    ws["A3"] = "Задолженность, млн руб."
    ws["B3"], ws["C3"], ws["D3"] = 10, 11, 12
    ws["A4"] = "Источник: техническое примечание"
    ws["B4"] = 2024
    wb.save(path)
    wb.close()

    sha = workbook_sha256(path)
    revision = "synthetic@sha256:" + sha
    raw = extract_raw_cells(path, source_id="synthetic", source_revision_id=revision, file_sha256=sha)
    observations, _, _, dispositions, diagnostics = parse_source_checked(
        path,
        spec=_spec(),
        source_revision_id=revision,
        file_sha256=sha,
        raw_cells=raw,
    )
    assert len(observations) == 3
    by_id = {row["raw_cell_id"]: row for row in dispositions}
    note_raw = next(row for row in raw if row["cell_coordinate"] == "B4")
    assert by_id[note_raw["raw_cell_id"]]["role"] == "source_metadata_numeric"
    assert diagnostics["numeric_disposition_gate"] == "passed"


def test_concept_identity_survives_row_movement_and_separates_measure_signature():
    spec = _spec()
    first = _concept(spec, sheet="Data", row=10, label="Задолженность", title="Таблица", row_axis="indicator")
    moved = _concept(spec, sheet="Data", row=11, label="Задолженность", title="Таблица", row_axis="indicator")
    assert first["source_concept_id"] == moved["source_concept_id"]

    regional = _spec("region")
    rub = _concept(
        regional,
        sheet="Debt",
        row=10,
        label="Москва",
        title="Задолженность, млн руб.",
        row_axis="region",
        unit="RUB",
        scale=1_000_000,
    )
    count = _concept(
        regional,
        sheet="Count",
        row=10,
        label="Москва",
        title="Количество кредитов, единиц",
        row_axis="region",
        unit="count",
        scale=1,
    )
    assert rub["source_concept_id"] != count["source_concept_id"]


def _validation_bundle():
    revision = "x@sha256:" + "a" * 64
    manifest = [{"source_id": "x", "source_revision_id": revision, "status": "ok"}]
    raw = [{
        "raw_cell_id": "rc_1",
        "source_id": "x",
        "source_revision_id": revision,
        "sheet_exact": "S",
        "cell_coordinate": "B2",
        "ooxml_type": "n",
        "value_lexical": "1",
        "formula": "",
    }]
    concepts = [{
        "source_concept_id": "sc_1",
        "source_id": "x",
        "source_local_key": "measure",
        "label_ru_source": "Показатель",
        "name_ru": "Показатель",
        "name_en": "Indicator",
        "translation_status": "project_translation",
    }]
    observations = [{
        "observation_id": "ob_1",
        "build_id": "bld_test",
        "source_id": "x",
        "source_revision_id": revision,
        "source_concept_id": "sc_1",
        "period": "2026-01-01",
        "frequency": "monthly",
        "period_role": "stock",
        "value_exact": "1",
        "unit": "RUB",
        "scale": "1",
        "dimensions_json": json.dumps({"a": "1", "b": "2"}),
        "sheet_exact": "S",
        "cell_coordinate": "B2",
        "raw_cell_id": "rc_1",
    }]
    dispositions = [{
        "raw_cell_id": "rc_1",
        "source_id": "x",
        "source_revision_id": revision,
        "role": "observation_value",
        "reason": "test",
    }]
    return manifest, raw, concepts, observations, dispositions


def _validate(bundle):
    manifest, raw, concepts, observations, dispositions = bundle
    return validate_bundle(
        manifest=manifest,
        raw_cells=raw,
        concepts=concepts,
        dimension_members=[],
        observations=observations,
        dispositions=dispositions,
        diagnostics=[{"source_id": "x"}],
        require_complete=False,
    )


def test_validation_blocks_cross_owner_source_mismatch_and_impossible_calendar():
    bundle = _validation_bundle()
    bundle[2][0]["source_id"] = "other"
    with pytest.raises(ValidationError, match="concept source ownership mismatch"):
        _validate(bundle)

    bundle = _validation_bundle()
    bundle[3][0]["period"] = "2026-99-99"
    with pytest.raises(ValidationError, match="Invalid ISO calendar period"):
        _validate(bundle)


def test_validation_canonicalizes_dimension_json_for_conflict_detection():
    manifest, raw, concepts, observations, dispositions = _validation_bundle()
    raw.append({**raw[0], "raw_cell_id": "rc_2", "cell_coordinate": "C2", "value_lexical": "2"})
    dispositions.append({**dispositions[0], "raw_cell_id": "rc_2"})
    observations.append({
        **observations[0],
        "observation_id": "ob_2",
        "raw_cell_id": "rc_2",
        "cell_coordinate": "C2",
        "value_exact": "2",
        "dimensions_json": json.dumps({"b": "2", "a": "1"}),
    })
    with pytest.raises(ValidationError, match="Conflicting semantic duplicates"):
        _validate((manifest, raw, concepts, observations, dispositions))


def test_cli_repeated_dimension_key_preserves_all_values():
    assert _dims(["region=Москва", "region=Санкт-Петербург"]) == {
        "region": ["Москва", "Санкт-Петербург"]
    }


def test_build_identity_binds_implementation_and_runtime(monkeypatch):
    spec = [get_source("mortgage_debt")]
    manifest = [{
        "source_id": "mortgage_debt",
        "source_revision_id": "mortgage_debt@sha256:" + "a" * 64,
        "status": "ok",
    }]
    import cbr_unified.build as build

    monkeypatch.setattr(build, "_implementation_fingerprint", lambda: ("a" * 64, {"python": "3.12.0"}))
    first, _ = _fingerprints(manifest, spec)
    monkeypatch.setattr(build, "_implementation_fingerprint", lambda: ("b" * 64, {"python": "3.12.0"}))
    changed_code, _ = _fingerprints(manifest, spec)
    monkeypatch.setattr(build, "_implementation_fingerprint", lambda: ("a" * 64, {"python": "3.12.1"}))
    changed_runtime, _ = _fingerprints(manifest, spec)
    assert first != changed_code
    assert first != changed_runtime
