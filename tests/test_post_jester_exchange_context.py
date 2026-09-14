from openpyxl import Workbook

from cbr_unified.processing import parse_source_checked
from cbr_unified.raw import extract_raw_cells, workbook_sha256
from cbr_unified.registry import SourceSpec


def _exchange_spec() -> SourceSpec:
    return SourceSpec(
        source_id="exchange_rate",
        url="https://cbr.ru/vfs/statistics/credit_statistics/ex_rate_ind/exchange_rate.xlsx",
        name_ru="Показатели валютного курса",
        name_en="Exchange-rate indicators",
        family="exchange_rates",
        row_axis="hierarchy",
        parser="exchange",
    )


def test_exchange_measure_headings_split_repeated_indicator_identity(tmp_path):
    path = tmp_path / "exchange-like.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Ежемесячные"
    ws["A1"] = "Показатели валютного курса"

    # Composite exchange calendar: one year heading spans monthly columns.
    ws["B2"] = 2024
    ws["B3"], ws["C3"], ws["D3"] = "январь", "февраль", "март"

    ws["B4"] = "Индексы обменного курса рубля к декабрю предыдущего года"
    ws["A5"] = "Реальный эффективный курс рубля"
    ws["B5"], ws["C5"], ws["D5"] = 101, 102, 103

    ws["B7"] = "Индексы обменного курса рубля к предыдущему периоду"
    ws["A8"] = "Реальный эффективный курс рубля"
    ws["B8"], ws["C8"], ws["D8"] = 99, 100, 101

    wb.save(path)
    wb.close()

    sha = workbook_sha256(path)
    revision = "exchange_rate@sha256:" + sha
    raw = extract_raw_cells(
        path,
        source_id="exchange_rate",
        source_revision_id=revision,
        file_sha256=sha,
    )
    observations, concepts, _, _, diagnostics = parse_source_checked(
        path,
        spec=_exchange_spec(),
        source_revision_id=revision,
        file_sha256=sha,
        raw_cells=raw,
    )

    first_ids = {
        row["source_concept_id"]
        for row in observations
        if row["source_row"] == 5
    }
    second_ids = {
        row["source_concept_id"]
        for row in observations
        if row["source_row"] == 8
    }

    assert len(observations) == 6
    assert len(first_ids) == 1
    assert len(second_ids) == 1
    assert first_ids != second_ids
    assert diagnostics["exchange_measure_context_rewrites"] == 6

    contexts = "\n".join(row["source_context"] for row in concepts)
    assert "к декабрю предыдущего года" in contexts
    assert "к предыдущему периоду" in contexts
