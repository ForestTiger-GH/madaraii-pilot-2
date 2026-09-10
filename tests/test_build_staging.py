from datetime import date

import pytest
from openpyxl import Workbook

from cbr_unified.build import build_database
from cbr_unified.registry import get_source


def _good_source(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "итого"
    ws["A1"] = "Задолженность по ипотечным жилищным кредитам, млн руб."
    ws["A2"] = "Регион"
    ws["B2"] = date(2026, 6, 1)
    ws["C2"] = date(2026, 7, 1)
    ws["D2"] = date(2026, 8, 1)
    ws["A3"] = "РОССИЙСКАЯ ФЕДЕРАЦИЯ"
    ws["B3"] = 1
    ws["C3"] = 2
    ws["D3"] = 3
    wb.save(path)


def _bad_source(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "unknown"
    ws["A1"] = "Неизвестная структура"
    for row in range(2, 8):
        for col in range(2, 6):
            ws.cell(row, col).value = row * col
    wb.save(path)


def test_successful_build_promotes_after_validation(tmp_path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    _good_source(input_dir / "mortgage_debt.xlsx")
    output = tmp_path / "product"
    output.mkdir()
    (output / "old.txt").write_text("previous", encoding="utf-8")

    result = build_database(
        output,
        input_dir=input_dir,
        sources=[get_source("mortgage_debt")],
        require_complete=False,
    )
    assert result.output_dir == output
    assert result.sqlite_path.exists()
    assert (output / "source_manifest.json").exists()
    assert not (output / "old.txt").exists()


def test_failed_build_keeps_previous_output_and_failure_evidence(tmp_path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    _bad_source(input_dir / "mortgage_debt.xlsx")
    output = tmp_path / "product"
    output.mkdir()
    sentinel = output / "old.txt"
    sentinel.write_text("previous", encoding="utf-8")

    with pytest.raises(RuntimeError, match="evidence retained"):
        build_database(
            output,
            input_dir=input_dir,
            sources=[get_source("mortgage_debt")],
            require_complete=False,
        )

    assert sentinel.read_text(encoding="utf-8") == "previous"
    staging = list(tmp_path.glob(".product.staging-*"))
    assert len(staging) == 1
    assert (staging[0] / "failure.json").exists()
    assert (staging[0] / "source_manifest.acquisition.json").exists()
    assert (staging[0] / "failure-evidence" / "mortgage_debt-raw.csv").exists()
