from datetime import date

from openpyxl import Workbook

from cbr_unified.raw import extract_raw_cells, workbook_sha256
from cbr_unified.registry import SourceSpec
from cbr_unified.semantic import parse_source


def _fixture_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "в рублях"
    ws["A1"] = "Задолженность тестовая в рублях, млн руб."
    ws["A2"] = "Показатель"
    ws["B2"] = date(2026, 6, 1)
    ws["C2"] = date(2026, 7, 1)
    ws["D2"] = date(2026, 8, 1)
    ws["A3"] = "ВСЕГО"
    ws["B3"] = 100
    ws["C3"] = 101.25
    ws["D3"] = 102
    wb.save(path)


def test_raw_extraction_and_semantic_lineage(tmp_path):
    path = tmp_path / "fixture.xlsx"
    _fixture_workbook(path)
    sha = workbook_sha256(path)
    revision = f"fixture@sha256:{sha}"
    raw = extract_raw_cells(path, source_id="fixture", source_revision_id=revision, file_sha256=sha)

    by_coord = {(r["sheet_exact"], r["cell_coordinate"]): r for r in raw}
    assert by_coord[("в рублях", "C3")]["value_lexical"] == "101.25"
    assert by_coord[("в рублях", "A3")]["text_resolved"] == "ВСЕГО"

    spec = SourceSpec(
        "fixture",
        "https://example.invalid/fixture.xlsx",
        "Тестовый показатель",
        "Test indicator",
        "test",
        "indicator",
        period_role="stock",
    )
    observations, concepts, members, dispositions, diagnostics = parse_source(
        path,
        spec=spec,
        source_revision_id=revision,
        file_sha256=sha,
        raw_cells=raw,
    )

    assert len(observations) == 3
    assert {o["period"] for o in observations} == {"2026-06-01", "2026-07-01", "2026-08-01"}
    assert {o["unit"] for o in observations} == {"RUB"}
    assert {o["scale"] for o in observations} == {1_000_000}
    assert all(o["source_revision_id"] == revision for o in observations)
    assert all(o["raw_cell_id"] in {r["raw_cell_id"] for r in raw} for o in observations)
    assert members == []
    assert concepts
    disposition_by_id = {d["raw_cell_id"]: d for d in dispositions}
    for obs in observations:
        assert disposition_by_id[obs["raw_cell_id"]]["role"] == "observation_value"
    assert diagnostics["observation_count"] == 3
