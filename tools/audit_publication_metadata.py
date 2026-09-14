from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


COMMENT_LIMIT = 100
FORMAT_LIMIT = 100
HIDDEN_LIMIT = 100


def _sample(items: list[dict[str, object]], limit: int) -> list[dict[str, object]]:
    return items[:limit]


def audit(source_dir: Path) -> dict[str, object]:
    files = sorted(source_dir.glob("*.xlsx"))
    hidden_sheets: list[dict[str, object]] = []
    hidden_rows: list[dict[str, object]] = []
    hidden_columns: list[dict[str, object]] = []
    comments: list[dict[str, object]] = []
    percent_formats: list[dict[str, object]] = []
    workbook_reports: list[dict[str, object]] = []

    for path in files:
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=True)
        try:
            wb_report = {
                "file": path.name,
                "sheet_count": len(wb.worksheets),
                "hidden_sheet_count": 0,
                "hidden_row_count": 0,
                "hidden_column_count": 0,
                "comment_count": 0,
                "numeric_percent_format_count": 0,
            }
            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    wb_report["hidden_sheet_count"] += 1
                    hidden_sheets.append({
                        "file": path.name,
                        "sheet": ws.title,
                        "state": ws.sheet_state,
                    })

                for row_index, dimension in ws.row_dimensions.items():
                    if dimension.hidden:
                        wb_report["hidden_row_count"] += 1
                        hidden_rows.append({
                            "file": path.name,
                            "sheet": ws.title,
                            "row": row_index,
                        })

                for column_key, dimension in ws.column_dimensions.items():
                    if dimension.hidden:
                        wb_report["hidden_column_count"] += 1
                        hidden_columns.append({
                            "file": path.name,
                            "sheet": ws.title,
                            "column": column_key,
                        })

                for row in ws.iter_rows():
                    for cell in row:
                        if cell.comment is not None:
                            wb_report["comment_count"] += 1
                            comments.append({
                                "file": path.name,
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "author": cell.comment.author or "",
                                "text": cell.comment.text or "",
                            })
                        value = cell.value
                        if (
                            isinstance(value, (int, float))
                            and not isinstance(value, bool)
                            and "%" in str(cell.number_format)
                        ):
                            wb_report["numeric_percent_format_count"] += 1
                            percent_formats.append({
                                "file": path.name,
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "value": value,
                                "number_format": cell.number_format,
                            })
            workbook_reports.append(wb_report)
        finally:
            wb.close()

    return {
        "source_file_count": len(files),
        "totals": {
            "hidden_sheets": len(hidden_sheets),
            "hidden_rows": len(hidden_rows),
            "hidden_columns": len(hidden_columns),
            "comments": len(comments),
            "numeric_percent_formats": len(percent_formats),
        },
        "samples": {
            "hidden_sheets": _sample(hidden_sheets, HIDDEN_LIMIT),
            "hidden_rows": _sample(hidden_rows, HIDDEN_LIMIT),
            "hidden_columns": _sample(hidden_columns, HIDDEN_LIMIT),
            "comments": _sample(comments, COMMENT_LIMIT),
            "numeric_percent_formats": _sample(percent_formats, FORMAT_LIMIT),
        },
        "workbooks": workbook_reports,
        "interpretation_boundary": (
            "This report establishes baseline applicability only. Presence of metadata does not by itself "
            "make it semantic, publication-controlling, or a Product defect; downstream engineering review owns that disposition."
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_dir")
    parser.add_argument("output")
    args = parser.parse_args()
    report = audit(Path(args.source_dir))
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["totals"], ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
