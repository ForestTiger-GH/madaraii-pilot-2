import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from cbr_unified.acquisition import _alternate_cbr_alias, _trusted_cbr_url, bind_local_sources
from cbr_unified.build import build_database
from cbr_unified.query import UnifiedDatabase
from cbr_unified.registry import SourceSpec, get_source


_AUDIT_TOOL = Path(__file__).resolve().parents[1] / "tools" / "audit_publication_metadata.py"
_AUDIT_SPEC = importlib.util.spec_from_file_location("audit_publication_metadata", _AUDIT_TOOL)
assert _AUDIT_SPEC is not None and _AUDIT_SPEC.loader is not None
_AUDIT_MODULE = importlib.util.module_from_spec(_AUDIT_SPEC)
_AUDIT_SPEC.loader.exec_module(_AUDIT_MODULE)
audit = _AUDIT_MODULE.audit


def _workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Задолженность, млн руб."
    ws["B2"], ws["C2"], ws["D2"] = "01.01.2024", "01.02.2024", "01.03.2024"
    ws["A3"] = "Москва"
    ws["B3"], ws["C3"], ws["D3"] = 10, 11, 12
    wb.save(path)
    wb.close()


def test_source_trust_rejects_non_cbr_hosts_and_malformed_local_xlsx(tmp_path):
    assert _trusted_cbr_url("https://cbr.ru/a.xlsx")
    assert _trusted_cbr_url("https://www.cbr.ru/a.xlsx")
    assert not _trusted_cbr_url("https://cbr.ru.attacker.example/a.xlsx")
    assert not _trusted_cbr_url("https://example.com/a.xlsx")
    assert _alternate_cbr_alias("https://www.cbr.ru/vfs/a.xlsx?x=1") == "https://cbr.ru/vfs/a.xlsx?x=1"
    assert _alternate_cbr_alias("https://cbr.ru/vfs/a.xlsx") is None
    assert _alternate_cbr_alias("https://sub.cbr.ru/vfs/a.xlsx") is None
    assert _alternate_cbr_alias("https://example.com/vfs/a.xlsx") is None

    bad = tmp_path / "payload.xlsx"
    bad.write_bytes(b"PK but not actually an OOXML workbook")
    records = bind_local_sources(
        tmp_path,
        sources=[get_source("mortgage_debt")],
        explicit_paths={"mortgage_debt": bad},
        require_complete=False,
    )
    assert records[0]["status"] == "failed"
    assert records[0]["error_type"] == "AcquisitionError"


def test_exact_build_sourcespec_is_used_by_semantic_parser(tmp_path):
    source_path = tmp_path / "mortgage_debt.xlsx"
    _workbook(source_path)
    original = get_source("mortgage_debt")
    custom = SourceSpec(
        source_id=original.source_id,
        url=original.url,
        name_ru=original.name_ru,
        name_en=original.name_en,
        family=original.family,
        row_axis=original.row_axis,
        parser=original.parser,
        classification=original.classification,
        period_role="flow",
    )
    result = build_database(
        tmp_path / "product",
        input_dir=tmp_path,
        sources=[custom],
        require_complete=False,
    )
    frame = UnifiedDatabase(result.sqlite_path).observations()
    assert set(frame["period_role"]) == {"flow"}


def test_publication_metadata_audit_detects_applicability_classes(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["C"].hidden = True
    ws["A1"] = "note"
    ws["A1"].comment = Comment("Предварительные данные. Не использовать для публикации.", "CBR")
    ws["B2"] = 0.25
    ws["B2"].number_format = "0.0%"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = 1
    wb.save(path)
    wb.close()

    report = audit(tmp_path)
    assert report["source_file_count"] == 1
    assert report["totals"]["hidden_sheets"] == 1
    assert report["totals"]["hidden_rows"] == 1
    assert report["totals"]["hidden_columns"] == 1
    assert report["totals"]["comments"] == 1
    assert report["totals"]["numeric_percent_formats"] == 1
