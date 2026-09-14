from openpyxl import Workbook

from cbr_unified.processing import parse_source_checked
from cbr_unified.raw import extract_raw_cells, workbook_sha256
from cbr_unified.semantic import _concept, _row_hierarchy_context
from cbr_unified.registry import SourceSpec


def _spec(row_axis="indicator", parser="matrix"):
    return SourceSpec(
        source_id="synthetic",
        url="https://cbr.ru/test.xlsx",
        name_ru="Тест",
        name_en="Test",
        family="test",
        row_axis=row_axis,
        parser=parser,
        period_role="stock",
    )


def test_repeated_child_labels_are_disambiguated_by_semantic_parent_not_row_number():
    wb = Workbook()
    ws = wb.active
    ws["A2"] = "Отчетная дата"
    ws["A3"] = "Количество кредитов, единиц, в том числе"
    ws["A4"] = "  на создание объектов, единиц"
    ws["A5"] = "  на приобретение объектов, единиц"
    ws["A6"] = "Объем кредитов, млн руб., в том числе"
    ws["A7"] = "  на создание объектов, млн руб."
    # Mirrors a real-source presentation typo: identical child label in another section.
    ws["A8"] = "  на приобретение объектов, единиц"

    first_context = _row_hierarchy_context(ws, 5, 2, 2)
    second_context = _row_hierarchy_context(ws, 8, 2, 2)
    assert first_context == "Количество кредитов, единиц, в том числе"
    assert second_context == "Объем кредитов, млн руб., в том числе"

    first = _concept(
        _spec(),
        sheet="Data",
        row=5,
        label="на приобретение объектов, единиц",
        title="Ипотечные кредиты",
        row_axis="indicator",
        unit="count",
        scale=1,
        identity_context=first_context,
    )
    second = _concept(
        _spec(),
        sheet="Data",
        row=8,
        label="на приобретение объектов, единиц",
        title="Ипотечные кредиты",
        row_axis="indicator",
        unit="count",
        scale=1,
        identity_context=second_context,
    )
    moved = _concept(
        _spec(),
        sheet="Data",
        row=50,
        label="на приобретение объектов, единиц",
        title="Ипотечные кредиты",
        row_axis="indicator",
        unit="count",
        scale=1,
        identity_context=first_context,
    )
    assert first["source_concept_id"] != second["source_concept_id"]
    assert first["source_concept_id"] == moved["source_concept_id"]
    wb.close()


def test_mixed_repeated_blocks_use_source_heading_not_physical_row(tmp_path):
    path = tmp_path / "mixed.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Сведения о заемщиках МСП"

    ws["A2"] = "Количество заемщиков, единиц"
    ws["B3"], ws["C3"], ws["D3"] = "01.01.2024", "01.02.2024", "01.03.2024"
    ws["A4"] = "Всего"
    ws["B4"], ws["C4"], ws["D4"] = 10, 11, 12

    ws["A6"] = "Объем кредитов, млн руб."
    ws["B7"], ws["C7"], ws["D7"] = "01.01.2024", "01.02.2024", "01.03.2024"
    ws["A8"] = "Всего"
    ws["B8"], ws["C8"], ws["D8"] = 20, 21, 22
    wb.save(path)
    wb.close()

    sha = workbook_sha256(path)
    revision = "synthetic@sha256:" + sha
    raw = extract_raw_cells(
        path,
        source_id="synthetic",
        source_revision_id=revision,
        file_sha256=sha,
    )
    observations, concepts, _, _, diagnostics = parse_source_checked(
        path,
        spec=_spec("mixed", "mixed"),
        source_revision_id=revision,
        file_sha256=sha,
        raw_cells=raw,
    )

    first_ids = {row["source_concept_id"] for row in observations if row["source_row"] == 4}
    second_ids = {row["source_concept_id"] for row in observations if row["source_row"] == 8}
    assert len(first_ids) == 1
    assert len(second_ids) == 1
    assert first_ids != second_ids
    assert diagnostics["period_block_concept_rewrites"] == 6
    assert len(concepts) == 2
    assert all("period_block" in row["source_local_key"] for row in concepts)
